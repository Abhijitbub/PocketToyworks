"""Retain GDD 1.2 formatting while replacing its two-mode policy with timed play."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document(root / 'Documentation/Pocket-Toyworks-GDD-1.2.docx')
def set_cell(cell, text):
    if cell.text == text:
        return
    paragraph = cell.paragraphs[0]
    props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.text = text
    if props is not None:
        paragraph.runs[0]._r.insert(0, props)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
updates = {
    'Version 1.2  |': 'Version 1.3  |  3 October 2026  |  Android first  |  Unity  |  Prototype 0.5.0',
    'This document defines': 'This document defines gameplay, content, repeat play, planned advertising, engineering and validation. Version 1.3 makes timed play the sole main-level flow, replacing the relaxed and challenge choices. The working title still requires store and trademark clearance before release.',
    'The main progression is relaxed': 'All main levels use a countdown and lives. The current offline prototype has a guided tutorial, ten levels, a toy shelf, profiles and an earned-coin Workshop. A future soft launch targets 60 levels and a small toy collection. The current build has no ads or in-app purchases. Multiplayer, paid currencies and a large workshop simulator remain outside the initial scope.',
    'Treat queues as': 'Treat queues as the main novelty to validate. Main levels use a timer and lives; the guided tutorial is untimed. Do not add color recipes, rotating conveyors or keys to the core puzzle.',
    'Use a deterministic solution': 'Use a deterministic solution for every authored level and a clear jam explanation. Undo is unlimited and free but never restores time. Timeout or restarting a started level costs one life. Home saves and pauses without a penalty. Future ads may appear only at defined breaks or explicit opt-in.',
    'Normal play and timed challenges': 'Timed play and lives',
    'Normal play is the default': 'Timed play is the default and only main-level flow. Play, Continue, replay and Next Level all use the countdown. Unlimited free undo supports spatial correction without replenishing time. A jam can be recovered through undo and never automatically removes a life. The guided tutorial remains untimed and exempt from life costs.',
    'The challenge clock begins': 'The clock begins when active gameplay first runs, even before placement. Undoing every move, receiving a time boost, visiting Home or replaying the tutorial does not create a new clock. Resetting a started attempt costs one life. Cancelled restarts leave the board and lives intact. Completed results are final; undo is available during an unfinished attempt.',
    'Starting or resuming a challenge': 'Starting or resuming an attempt does not deduct a life. Timeout or confirmed restart is the charge point. At zero lives, new attempts are blocked; saved unfinished attempts can resume. Offer a coin refill, free regeneration and Workshop gifts. Home, profiles, the shelf and tutorial remain accessible. Never offer an untimed main-level fallback.',
    'Initial challenge limits': 'Initial limits are 3 minutes for levels 1 through 5 and 4 minutes for levels 6 through 10. Time boosts add 30 seconds and unlock at level 4. These are prototype tuning values; player testing determines later difficulty.',
    'The main screen reserves': 'The main screen places a visible countdown and level number at the top, mould grids in the middle, then reserves and three supply lanes. Lives, coins and Learn sit below the timer. Tap-select plus tap-place is an alternative to dragging. Selecting a piece reveals Rotate and its ghost footprint. Home level cards have one Play or Continue button and show the time limit.',
    'The current profile provides': 'The profile has eight original toy characters and six frames. Two frames are initially owned; the others cost 100, 150, 200 and 250 earned coins. A daily gift grants 75 coins every 24 hours. Two time boosts cost 50 coins, an auto-fit costs 80, and a full life refill costs 150. Time boosts unlock at level 4 and auto-fit at level 7. Lives cap at five and regenerate every 30 minutes. Palette customization remains future work.',
    'Repeat play without forced pressure': 'Repeat play and rewards',
    'The prototype shelf displays': 'The prototype shelf displays completed orders and locked future orders, with replay access and original toy portraits. Profile characters and frames are available now. Additional toy families, palettes and freely arranged shelf sockets remain future expansion. Cosmetic ownership never changes the timer or life rules.',
    'On return, resume': 'On return, resume the exact puzzle or open the shelf in one tap. Home, settings, backgrounding and tutorial replay pause the clock without using a life. Switching levels preserves their separate boards and timers. Notifications remain optional. Taking a break must not reset the board or cost a life.',
    'Save the active puzzle': 'Save active and paused attempts by level ID, including history, elapsed play, remaining time and clock-start status. On upgrade, retain boards, profiles, rewards and balances. An existing timed attempt keeps its remaining time; an untimed 0.4.0 attempt receives a full clock once. Tutorial suspension follows the same migration. Where old paused modes share a level, prefer the timed attempt. Regenerate one life every 30 minutes, up to five, with clock-rollback protection.',
    'The queue is a stock': 'The queue is a stock of future pieces; selectable fronts constrain actions. Reserves provide limited working capacity. Placement consumes empty space while mould completion supplies new space. The timer is a resource undo cannot replenish. Coins have explicit sources and sinks; cosmetic ownership does not amplify earning power. [9, chs 2, 7, 8, 10]',
    'Normal play, undo, hints': 'Timed puzzles remain available offline. At zero lives, preserve saves and offer regeneration or earned-coin refills. Tutorial and Home remain accessible. Never require an ad to repair an unsolvable level or alter queues according to spending or ad behavior.',
}
for paragraph in doc.paragraphs:
    for prefix, replacement in updates.items():
        if paragraph.text.startswith(prefix):
            paragraph.text = replacement
            break

for table in doc.tables:
    if table.cell(0, 0).text == 'Action':
        values = [
            ('Action', 'Main levels', 'Guided tutorial'),
            ('Undo', 'Free; never restores time', 'Free'),
            ('Restart', 'One life after the clock starts', 'Free'),
            ('Home or pause', 'Save and pause for free', 'Restore the suspended level'),
            ('Failure', 'Timeout removes one life once', 'No timer or life penalty'),
            ('Zero lives', 'Resume saved attempts; no new starts', 'Available for learning'),
        ]
        for row, texts in zip(table.rows, values):
            for cell, text in zip(row.cells, texts):
                set_cell(cell, text)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            text = cell.text
            if text.startswith('Normal play has unlimited'):
                set_cell(cell, 'Undo is free and never restores time. Restart after the clock starts or timeout uses one life. Home saves and pauses for free.')
            elif text.startswith('A future opt-in rewarded ad may add'):
                set_cell(cell, 'A future opt-in rewarded ad may add a third reserve for an attempt. Baseline puzzles remain solvable with two. Restart keeps the ordinary life rule.')
            elif text.startswith('Once per attempt, available after level 10.'):
                set_cell(cell, 'Once per attempt after level 10. This is future work; the current build has two reserves.')
            elif text.startswith('Use original layout and assets.'):
                set_cell(cell, 'Use original layout and assets. Door matching and color restrictions are outside our rules. All main levels use timed play.')
            else:
                set_cell(cell, text.replace('The normal Continue button', 'The Continue button').replace('paused mode sessions', 'paused level attempts').replace('challenge lives and clocks', 'lives and clocks'))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if ri == 0:
                        run.bold = True

out = root / 'Documentation/Pocket-Toyworks-GDD-1.3.docx'
doc.save(out)
(root / 'Documentation/GDD-current.txt').write_text('\n'.join([p.text for p in doc.paragraphs] + ['\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows) for t in doc.tables]), encoding='utf-8')
print(out)
