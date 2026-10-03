"""Apply the approved normal/challenge policy to the retained Word GDD."""
from pathlib import Path
from copy import deepcopy
from shutil import copy2
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).resolve().parents[1]
source = Path(r'C:/Users/abzzz/OneDrive/Desktop/GDD/Pocket-Toyworks-GDD.docx')
original = root / 'Documentation/Pocket-Toyworks-GDD-1.1-original.docx'
if not original.exists():
    copy2(source, original)
doc = Document(original)
p = doc.paragraphs
tables = list(doc.tables)
updates = {
    0: 'Pocket Toyworks game design',
    1: 'Gameplay production and validation',
    2: 'Version 1.2  |  2 October 2026  |  Android first  |  Unity  |  Prototype 0.4.0',
    4: 'This document defines the gameplay, content pipeline, repeat-play structure, planned advertising, acquisition experiments, engineering boundaries, and decision gates. Version 1.2 separates free relaxed puzzles from optional timed challenges. The working title still requires store and trademark clearance before release.',
    10: 'The main progression is relaxed play with optional timed challenges on the same authored puzzles. The current offline prototype has a guided tutorial, ten levels, a toy shelf, player profiles and an earned-coin Workshop. A future soft launch targets 60 levels and a small toy collection. Advertising remains a future business-model experiment; the current build has no ads or in-app purchases. Multiplayer, paid currencies and a large workshop simulator remain outside the initial scope.',
    22: 'Treat queues as the main novelty to validate. Keep the timer and lives confined to optional challenges; do not add color recipes, rotating conveyors or keys to the core puzzle.',
    31: 'Use a deterministic solution for every authored level, unlimited free undo and restart in normal play, and a clear jam explanation. Timed challenges keep undo free without restoring time; lives apply only to challenge failure or resetting a started challenge. Home saves progress without a penalty. Future ads may appear only at defined breaks or explicit opt-in.',
    42: 'The main screen reserves the top area for order progress, the middle for large mould grids, and the lower area for reserve slots and three supply lanes. Settings and undo sit outside the drag area. Tap-select plus tap-place is an alternative to dragging. Selecting a piece reveals Rotate and a ghost footprint. Normal play shows a Free play indicator; the timer and life counter appear in challenge gameplay.',
    49: 'A completed main level unlocks the next and earns 50 workshop coins on its first clear only. The prototype has ten authored levels and a shelf entry for each completed order, using the original Pip toy reveal. Repeated clears and undo cannot duplicate rewards. The later 60-level content plan targets 12 reusable toy model families; this expansion requires the content and retention gates below.',
    50: 'The current profile provides eight original toy characters and six frames. Two frames are initially owned; the other four cost 100, 150, 200 and 250 earned coins. A daily gift grants 75 coins once every 24 hours. Two time boosts cost 50 coins, one verified auto-fit costs 80 coins, and a full challenge-life refill costs 150 coins. Time boosts unlock at level 4 and auto-fit at level 7. Normal play is always available without lives. Palette customization remains a future cosmetic feature.',
    51: 'A later daily workshop puzzle may unlock after level 15, using a curated, prevalidated 14-puzzle pool separate from the planned 60 main levels. Rotate by UTC date and keep seven past dates accessible. A dated daily ID may grant a stamp and five workshop coins once. There are no broken-streak penalties, and calendar refresh must not change an active board. This daily puzzle pool is not implemented in the current prototype.',
    52: 'On return, resume the exact puzzle or open the shelf in one tap. Home, settings, backgrounding and the tutorial pause a challenge clock without consuming a life. Switching modes preserves a separate puzzle and remaining time for each level and mode. Notifications remain optional. Return should be motivated by mastery and collection, not penalties for taking a break.',
    66: 'Mould reveals are provisional until the level is won. Persist board snapshots after every committed action. Commit unlocks and workshop coins atomically once per completed level ID. Undo can reverse a reveal but cannot duplicate permanent rewards or challenge-win credit. Keep challenge time, attempt-start status, life balance and owned tools outside puzzle undo snapshots.',
    67: 'Save the active puzzle and paused attempts by level ID and mode, including move history, elapsed play time, remaining challenge time and whether its clock has started. Migrate existing saves without changing their boards, rewards or life balance. A legacy active challenge is treated as started. Regenerate one challenge life every 30 minutes up to five, with clock-rollback protection. Future extra-reserve assistance remains outside this build.',
    72: 'The queue is a stock of future pieces; selectable fronts constrain actions. Reserve slots are limited working capacity. Placement consumes empty space, while mould completion supplies new space. Normal play supports free experimentation. Optional challenges add a time resource that undo cannot replenish. Workshop coins have explicit sources and sinks; cosmetic ownership does not amplify earning power. [9, chs 2, 7, 8, 10]',
    74: 'Future advertising design',
    75: 'This build has no ads and all hints are free. The placements below are future proposals. Use natural breaks for interstitials and a clear exchange for rewarded ads; verify SDK and store requirements before implementation. [10, 11]',
    81: 'Normal play, undo, hints, restart and the tutorial remain available offline and at zero challenge lives. Never require an ad to repair an unsolvable level or alter queues according to spending or ad behavior.',
    115: 'On completion, briefly lift the filled mould and assemble it into the toy. Intermediate reveals stay within gameplay; the final reveal leads to a result screen. Keep Continue prominent. First-clear coins commit once and completed orders remain on the shelf. Optional advertising bonuses remain future work.',
    117: 'The prototype shelf displays completed orders and locked future orders, with replay access and original toy portraits. Profile characters and frames are available now. Additional toy families, palettes and freely arranged shelf sockets remain future expansion; they must not introduce a life requirement for ordinary play.',
}
for index, text in updates.items():
    p[index].text = text

anchor = p[33]
def insert(text, style='Normal'):
    return anchor.insert_paragraph_before(text, style)
insert('Normal play and timed challenges', 'Heading 1').paragraph_format.page_break_before = True
insert('Normal play is the default, untimed experience. Unlimited undo and restart support spatial experimentation. Lives have a distinct purpose in optional timed challenges, where resetting the timer changes the attempt. A jam is recoverable through undo and never automatically removes a life.')
table = doc.add_table(rows=1, cols=3)
table.style = tables[2].style
table.autofit = False
for column, width in zip(table.columns, [1.1, 2.1, 3.1]):
    column.width = Inches(width)
for cell, text, width in zip(table.rows[0].cells, ['Action', 'Normal play', 'Timed challenge'], [1.1, 2.1, 3.1]):
    cell.text = text
    cell.width = Inches(width)
for values in [
    ('Undo', 'Free and unlimited', 'Free; remaining time is not restored'),
    ('Restart', 'Free', 'One life after the clock starts; confirmation first'),
    ('Home or pause', 'Save the puzzle for free', 'Save the puzzle and pause the clock for free'),
    ('Failure', 'No time limit or life loss', 'Timeout removes exactly one life'),
    ('Zero lives', 'Play normally and use the tutorial', 'Resume a saved attempt; new attempts need a life'),
]:
    for cell, text in zip(table.add_row().cells, values):
        cell.text = text
anchor._p.addprevious(table._tbl)
insert('The challenge clock begins when active gameplay first runs, even before a placement. Undoing to the initial board, receiving a time boost, opening Home or replaying the tutorial does not create a new clock. Starting at full time after a started attempt is Restart and costs one life. Completed and already failed attempts are not charged again. Undo after a completed challenge returns to free practice.').paragraph_format.space_before = Pt(7)
insert('Starting or resuming a challenge does not deduct a life. Timeout or a confirmed restart is the charge point. Once the last life is used on restart, show Challenge lives and offer Play relaxed. Preserve any separately paused normal puzzle. A cancelled restart changes neither the board nor the life count.')
insert('Initial challenge limits are 3 minutes for levels 1 through 5 and 4 minutes for levels 6 through 10. Time boosts add 30 seconds and apply only to timed play. These are tuning values for the prototype; player testing determines later difficulty settings.')

tables[2].rows[-1].cells[1].text = 'Normal play has unlimited free undo and restart, no timer and no life requirement. Undo restores queue heads, reserves and mould replacement. Optional timed challenges follow the separate mode policy above.'
tables[6].rows[2].cells[1].text = 'A future opt-in rewarded ad may add one third reserve slot for an attempt. Baseline puzzles remain solvable with two. Normal restart stays free.'
tables[6].rows[3].cells[1].text = 'A future optional rewarded ad after a first clear may add 10 workshop coins, without changing toy unlocks.'
tables[10].rows[1].cells[2].text = 'Use original layout and assets. Door matching and color restrictions are outside our rules. Timers belong only to optional challenges.'
tables[11].rows[1].cells[1].text = 'Pale blue gameplay background #D8E7FF; dark blue panel fill #2E66C3; cyan bevels and dark outlines. Piece accents: coral, yellow, blue and violet. Home uses a blue to violet backdrop. Review contrast on phones.'
tables[11].rows[4].cells[1].text = 'Use scalable TextMeshPro distance-field display and body fonts, clear outlines, safe-area margins and generous touch targets. Keep lane headings wholly inside the dark blue fill. Review portrait 16:9 and 18:9, native pixels, empty states and popups before delivery.'
tables[13].rows[4].cells[1].text = 'Atomic completion, coins, collection, profiles, challenge lives and clocks, paused mode sessions, migration and interrupted-write recovery.'

# Keep the retained document typography, with explicit black headings and clean titles.
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
for paragraph in doc.paragraphs:
    if paragraph.style.name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
        borders = paragraph._p.find('w:pPr/w:pBdr', paragraph._p.nsmap)
        if borders is not None:
            borders.getparent().remove(borders)
for t in doc.tables:
    borders = t._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders'); t._tbl.tblPr.append(borders)
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = borders.find(qn('w:' + edge))
        if e is None:
            e = OxmlElement('w:' + edge); borders.append(e)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9')
    for ri, row in enumerate(t.rows):
        # Repeating headers and flowing rows prevent clipping at page breaks.
        if ri == 0:
            h = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(h)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shd = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd'); cell._tc.get_or_add_tcPr().append(shd)
            shd.set(qn('w:fill'), 'E6EDF4' if ri == 0 else 'FFFFFF' if ri % 2 else 'F5F7FA')
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if ri == 0: run.bold = True
out = root / 'Documentation/Pocket-Toyworks-GDD-1.2.docx'
doc.save(out)
(root / 'Documentation/GDD-current.txt').write_text('\n'.join([p.text for p in doc.paragraphs] + ['\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows) for t in doc.tables]), encoding='utf-8')
print(out)
